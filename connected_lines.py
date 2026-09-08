# -*- coding: utf-8 -*-
import os, sys, glob, re
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def decrypt(buf):
    return bytes([(b - i - len(buf)) & 0xFF for i, b in enumerate(buf)])

def extract_strings_with_gaps(g2_path):
    with open(g2_path, 'rb') as f:
        data = f.read()
    strings = []
    i = 0
    data_len = len(data)
    while i < data_len - 6:
        tag = data[i] | (data[i+1] << 8)
        length = data[i+2] | (data[i+3] << 8) | (data[i+4] << 16) | (data[i+5] << 24)
        if tag in (0x0100, 0x0104, 0x0200) and length < 384 and i + 6 + length <= data_len:
            chunk = data[i+6 : i+6+length]
            d = decrypt(chunk)
            try:
                s = d.decode('cp932')
                strings.append({
                    'offset': i,
                    'end': i + 6 + length,
                    'text': s,
                })
                i += 6 + length
                continue
            except:
                pass
        i += 1
    return strings, data

def find_connected_groups(strings, data):
    groups = []
    curr = []
    for k in range(len(strings) - 1):
        cur = strings[k]
        nxt = strings[k+1]
        gap = data[cur['end'] : nxt['offset']]
        is_conn = (len(gap) == 10 and gap[0] == 0x21 and gap[5:] == b'\x06\x05\x01\x10\x00')
        if is_conn:
            if not curr:
                curr.append(k)
            curr.append(k + 1)
        else:
            if curr:
                groups.append(curr)
                curr = []
    if curr:
        groups.append(curr)
    return groups

LINE_LIMIT = 45  # Строгий лимит символов на одну строку экрана

def clean_text(t):
    if not t:
        return ''
    t = str(t).replace('\\n', ' ').replace('\n', ' ').replace('\r', ' ')
    t = re.sub(r'[-–—#]\s+', '', t)
    t = re.sub(r'\s+', ' ', t).strip()
    return t

def strict_split(text, num_lines, limit=LINE_LIMIT):
    """
    Строго разбивает текст на num_lines частей по границам целых слов.
    Первые (num_lines - 1) строк ГАРАНТИРОВАННО <= limit символов (без переносов внутри!).
    Последняя строка берет остаток текста.
    """
    words = clean_text(text).split(' ')
    words = [w for w in words if w]
    if not words:
        return [''] * num_lines

    lines = []
    current = []

    for word in words:
        cand = (' '.join(current + [word])).strip()
        if len(cand) <= limit or not current:
            current.append(word)
        else:
            if len(lines) < num_lines - 1:
                lines.append(' '.join(current))
                current = [word]
            else:
                current.append(word)

    if current:
        lines.append(' '.join(current))

    while len(lines) < num_lines:
        lines.append('')

    return lines

def cmd_audit(orig_dir, xlsx_path, out_audit_path):
    orig_path = Path(orig_dir)
    if not orig_path.is_dir():
        print(f'Папка {orig_dir} не найдена!')
        return

    print(f'Загрузка {xlsx_path}...')
    wb_in = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = 'ConnectedLines'

    headers = [
        'Sheet', 'Type', 'Rows', 'Status',
        'RU Full (Editable)', 'Total Len',
        'Suggested L1', 'Len1',
        'Suggested L2', 'Len2',
        'Suggested L3', 'Len3',
        'Current L1', 'CurLen1',
        'Current L2', 'CurLen2',
        'JP Full'
    ]
    ws_out.append(headers)

    hdr_fill = PatternFill('solid', fgColor='2F4F8F')
    hdr_font = Font(bold=True, color='FFFFFF')
    for col_idx in range(1, len(headers) + 1):
        cell = ws_out.cell(row=1, column=col_idx)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    warn_fill = PatternFill('solid', fgColor='FFC7CE')
    warn_font = Font(color='9C0006', bold=True)
    yellow_fill = PatternFill('solid', fgColor='FFF2CC')
    yellow_font = Font(color='7F6000', bold=True)

    total_chains = 0
    cur_overflow_l1 = 0
    too_long_count = 0

    for sheet_name in wb_in.sheetnames:
        g2_p = orig_path / f'{sheet_name}.g2'
        if not g2_p.exists():
            continue

        strings, data = extract_strings_with_gaps(g2_p)
        groups = find_connected_groups(strings, data)
        if not groups:
            continue

        ws = wb_in[sheet_name]
        row_map = {}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or r[0] is None:
                continue
            try:
                rid = int(r[0])
                jp = str(r[1]) if len(r) > 1 and r[1] is not None else ''
                ru = str(r[2]) if len(r) > 2 and r[2] is not None else ''
                row_map[rid] = (jp, ru)
            except (ValueError, TypeError):
                continue

        for grp in groups:
            chain_len = len(grp)
            rows_str = ', '.join(str(idx) for idx in grp)
            
            jp_parts = []
            ru_parts = []
            for idx in grp:
                jp, ru = row_map.get(idx, ('', ''))
                jp_parts.append(jp)
                ru_parts.append(ru)

            jp_full = ''.join(jp_parts)
            ru_full = clean_text(' '.join([clean_text(p) for p in ru_parts if p.strip()]))

            cur_l1 = clean_text(ru_parts[0])
            cur_l2 = clean_text(ru_parts[1]) if len(ru_parts) > 1 else ''
            cur_len1 = len(cur_l1)
            cur_len2 = len(cur_l2)

            suggested = strict_split(ru_full, chain_len, LINE_LIMIT)
            while len(suggested) < 3:
                suggested.append('')

            len_sug1 = len(suggested[0])
            len_sug2 = len(suggested[1])
            len_sug3 = len(suggested[2])
            tot_len = len(ru_full)

            max_capacity = chain_len * LINE_LIMIT
            issues = []
            if cur_len1 > LINE_LIMIT:
                issues.append('CUR_L1_OVERFLOW')
                cur_overflow_l1 += 1
            if tot_len > max_capacity:
                issues.append(f'TOO_LONG(>{max_capacity})')
                too_long_count += 1
            
            status = ' | '.join(issues) if issues else 'OK'

            row_data = [
                sheet_name, f'{chain_len}-lines', rows_str, status,
                ru_full, tot_len,
                suggested[0], len_sug1,
                suggested[1], len_sug2,
                suggested[2] if chain_len == 3 else '', len_sug3 if chain_len == 3 else '',
                cur_l1, cur_len1,
                cur_l2, cur_len2,
                jp_full
            ]
            ws_out.append(row_data)
            total_chains += 1

            curr_row = ws_out.max_row
            status_cell = ws_out.cell(row=curr_row, column=4)
            if 'CUR_L1_OVERFLOW' in status:
                status_cell.fill = warn_fill
                status_cell.font = warn_font
            elif 'TOO_LONG' in status:
                status_cell.fill = yellow_fill
                status_cell.font = yellow_font

    col_widths = {
        'A': 10, 'B': 10, 'C': 12, 'D': 22,
        'E': 55, 'F': 10,
        'G': 45, 'H': 8,
        'I': 45, 'J': 8,
        'K': 40, 'L': 8,
        'M': 45, 'N': 8,
        'O': 45, 'P': 8,
        'Q': 45
    }
    for col_letter, w in col_widths.items():
        ws_out.column_dimensions[col_letter].width = w

    ws_out.freeze_panes = 'A2'
    ws_out.auto_filter.ref = f'A1:Q{ws_out.max_row}'
    
    wb_out.save(out_audit_path)
    print(f'Готово! Всего связанных реплик: {total_chains}')
    print(f'Из них в aster.xlsx L1 превышает 45 символов (CUR_L1_OVERFLOW): {cur_overflow_l1}')
    print(f'Суммарно не влезает в 2/3 строки (TOO_LONG > 90/135): {too_long_count}')
    print(f'Сохранено в: {out_audit_path}')

def cmd_apply(audit_path, in_xlsx_path, out_xlsx_path):
    print(f'Загрузка audit: {audit_path}...')
    wb_audit = openpyxl.load_workbook(audit_path, data_only=True)
    ws_audit = wb_audit.active

    print(f'Загрузка исходной таблицы: {in_xlsx_path}...')
    wb_target = openpyxl.load_workbook(in_xlsx_path)

    print('Индексация целевой таблицы...')
    sheet_row_maps = {}
    for sheet_name in wb_target.sheetnames:
        tws = wb_target[sheet_name]
        rmap = {}
        for r in tws.iter_rows(min_row=2):
            if r[0].value is not None:
                try:
                    rmap[int(r[0].value)] = r
                except (ValueError, TypeError):
                    pass
        sheet_row_maps[sheet_name] = rmap

    updates = 0
    for row in ws_audit.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        sheet_name = str(row[0]).strip()
        rows_str = str(row[2]).strip()
        
        # Берем Suggested Line 1 (col G = index 6), Line 2 (col I = index 8), Line 3 (col K = index 10)
        sug_1 = str(row[6]).strip() if len(row) > 6 and row[6] else ''
        sug_2 = str(row[8]).strip() if len(row) > 8 and row[8] else ''
        sug_3 = str(row[10]).strip() if len(row) > 10 and row[10] else ''

        if sheet_name not in sheet_row_maps:
            continue

        row_indices = [int(x.strip()) for x in rows_str.split(',') if x.strip().isdigit()]
        if not row_indices:
            continue

        line_texts = [sug_1, sug_2, sug_3][:len(row_indices)]

        rmap = sheet_row_maps[sheet_name]
        for target_rid, new_text in zip(row_indices, line_texts):
            if target_rid in rmap:
                row_cells = rmap[target_rid]
                row_cells[2].value = new_text
                updates += 1

    wb_target.save(out_xlsx_path)
    print(f'Готово! Обновлено строк: {updates} -> {out_xlsx_path}')

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Использование:')
        print('  python connected_lines.py audit <ORIG> <aster.xlsx> <audit.xlsx>')
        print('  python connected_lines.py apply <audit.xlsx> <aster.xlsx> [out_aster.xlsx]')
        sys.exit(1)

    mode = sys.argv[1].lower()
    if mode == 'audit' and len(sys.argv) >= 5:
        cmd_audit(sys.argv[2], sys.argv[3], sys.argv[4])
    elif mode == 'apply' and len(sys.argv) >= 4:
        out_p = sys.argv[4] if len(sys.argv) >= 5 else sys.argv[3]
        cmd_apply(sys.argv[2], sys.argv[3], out_p)
    else:
        print('Неверные параметры')
